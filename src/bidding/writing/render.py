"""Markdown assembly and styled Word export for generated bid drafts."""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from ..models import now_iso
from ..parsing.structure import normalize_ws
from .outline import BidOutline, BidSection
from .profile import (
    BidderProfile,
    COMPANY_FIELDS,
    CREDIT_FIELDS,
    FINANCIAL_FIELDS,
    MANAGEMENT_SYSTEM_FIELDS,
    PERSONNEL_FIELDS,
    QUALIFICATION_FIELDS,
    SAFETY_FIELDS,
    has_placeholder,
    placeholder,
)
from .tender import TenderContext

PENDING_RE = re.compile(r"【待(?:填写|核对)(?:[：:][^】]*)?】")
_TABLE_SEPARATOR_RE = re.compile(r"^:?-{3,}:?$")
_HEADING_RE = re.compile(r"^(#{1,4})\s+(.*)$")
_BULLET_RE = re.compile(r"^\s*[-*+]\s+(.*)$")
_NUMBERED_RE = re.compile(r"^\s*(\d+)[.)、]\s+(.*)$")
_INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")
_CN_DIGITS = "零一二三四五六七八九"


@dataclass(slots=True)
class PendingMarker:
    marker: str
    sections: list[str]


def cn_index(number: int) -> str:
    if number <= 0:
        return str(number)
    if number < 10:
        return _CN_DIGITS[number]
    if number == 10:
        return "十"
    if number < 20:
        return "十" + _CN_DIGITS[number % 10]
    if number < 100:
        tens, ones = divmod(number, 10)
        return _CN_DIGITS[tens] + "十" + (_CN_DIGITS[ones] if ones else "")
    return str(number)


def _cell(value: Any) -> str:
    text = str(value or "").strip()
    if not text:
        return ""
    return text.replace("|", "\\|").replace("\r\n", "<br>").replace("\n", "<br>")


def _md_table(headers: Iterable[str], rows: Iterable[Iterable[Any]]) -> str:
    header_line = "| " + " | ".join(_cell(item) for item in headers) + " |"
    separator = "| " + " | ".join("---" for _ in headers) + " |"
    body = [
        "| " + " | ".join(_cell(item) for item in row) + " |"
        for row in rows
    ]
    return "\n".join([header_line, separator, *body])


def _kv_table(pairs: Iterable[tuple[str, str]]) -> str:
    return _md_table(("项目", "内容"), [(label, value) for label, value in pairs])


def _project_manager(profile: BidderProfile) -> str:
    for row in profile.personnel:
        role = row.get("role", "")
        name = row.get("name", "")
        if "项目经理" in role and name and not has_placeholder(name):
            return name
    return placeholder("项目经理姓名")


def render_profile_table(profile: BidderProfile, table_id: str) -> str:
    if table_id == "company_basic":
        pairs = [
            (label, profile.company.get(key, "")) for key, label in COMPANY_FIELDS.items()
        ]
        return "**附表：投标人基本情况**\n\n" + _kv_table(pairs)
    if table_id == "company_identity":
        pairs = [
            ("投标人名称", profile.company.get("name", "")),
            ("统一社会信用代码", profile.company.get("unified_social_credit_code", "")),
            ("法定代表人", profile.company.get("legal_representative", "")),
            ("注册地址", profile.company.get("address", "")),
            ("成立日期", profile.company.get("established_date", "")),
            ("联系人", profile.company.get("contact_person", "")),
            ("联系电话", profile.company.get("phone", "")),
        ]
        return "**附表：签署主体信息**\n\n" + _kv_table(pairs)
    if table_id == "qualifications":
        rows = [
            [
                index,
                row.get("name", ""),
                row.get("level", ""),
                row.get("certificate_no", ""),
                row.get("issuer", ""),
                row.get("valid_until", ""),
                row.get("note", ""),
            ]
            for index, row in enumerate(profile.qualifications, start=1)
        ]
        return "**附表：资质证书**\n\n" + _md_table(
            ("序号", "资质名称", "等级", "证书编号", "发证机关", "有效期至", "备注"), rows
        )
    if table_id == "credit":
        pairs = [(label, profile.credit.get(key, "")) for key, label in CREDIT_FIELDS.items()]
        return "**附表：信用状况**\n\n" + _kv_table(pairs)
    if table_id == "management_systems":
        rows = [
            [
                index,
                row.get("name", ""),
                row.get("standard", ""),
                row.get("certificate_no", ""),
                row.get("valid_until", ""),
            ]
            for index, row in enumerate(profile.management_systems, start=1)
        ]
        return "**附表：管理体系认证**\n\n" + _md_table(
            ("序号", "体系名称", "认证标准", "证书编号", "有效期至"), rows
        )
    if table_id == "safety":
        pairs = [(label, profile.safety.get(key, "")) for key, label in SAFETY_FIELDS.items()]
        return "**附表：安全生产信息**\n\n" + _kv_table(pairs)
    if table_id == "financials":
        rows = [
            [
                row.get("year", ""),
                row.get("revenue", ""),
                row.get("net_profit", ""),
                row.get("total_assets", ""),
                row.get("total_liabilities", ""),
                row.get("audit_status", ""),
            ]
            for row in profile.financials
        ]
        return "**附表：近年财务状况**\n\n" + _md_table(
            ("年度", "营业收入", "净利润", "总资产", "总负债", "审计情况"), rows
        )
    if table_id == "similar_projects":
        rows = []
        for index, row in enumerate(profile.similar_projects, start=1):
            period = f"{row.get('start_date', '')} 至 {row.get('end_date', '')}"
            role_text = row.get("role", "")
            if row.get("description"):
                role_text = f"{role_text}；{row['description']}"
            rows.append(
                [
                    index,
                    row.get("project_name", ""),
                    row.get("owner", ""),
                    row.get("contract_amount", ""),
                    period,
                    role_text,
                    row.get("proof", ""),
                ]
            )
        return "**附表：类似项目业绩**\n\n" + _md_table(
            ("序号", "项目名称", "发包人", "合同金额", "开竣工日期", "角色及主要工作", "证明材料"),
            rows,
        )
    if table_id == "personnel":
        rows = []
        for index, row in enumerate(profile.personnel, start=1):
            qualification = "；".join(
                value
                for value in (row.get("title", ""), row.get("certificate", ""))
                if value
            )
            rows.append(
                [
                    index,
                    row.get("name", ""),
                    row.get("role", ""),
                    qualification,
                    row.get("certificate_no", ""),
                    row.get("experience_years", ""),
                    row.get("projects", ""),
                ]
            )
        return "**附表：项目管理机构及人员**\n\n" + _md_table(
            ("序号", "姓名", "拟任岗位", "职称/执业资格", "证书编号", "从业年限", "类似项目经历"),
            rows,
        )
    if table_id == "bid_letter_appendix":
        decision_rows = [
            ("投标总报价（小写）", placeholder("投标总报价（小写）")),
            ("投标总报价（大写）", placeholder("投标总报价（大写）")),
            ("工期", placeholder("投标工期")),
            ("工程质量", placeholder("工程质量承诺")),
            ("投标有效期", placeholder("投标有效期")),
            ("投标保证金", placeholder("投标保证金金额及形式")),
            ("项目经理", _project_manager(profile)),
            ("缺陷责任期", placeholder("缺陷责任期")),
            ("分包情况", placeholder("分包情况")),
            ("违约金", placeholder("违约金金额")),
            ("工程预付款", placeholder("工程预付款")),
            ("质量保证金", placeholder("质量保证金")),
        ]
        rows = [[index, label, value] for index, (label, value) in enumerate(decision_rows, start=1)]
        return "**附表：投标函附录填写项**\n\n" + _md_table(
            ("序号", "填写项", "内容"), rows
        )
    if table_id == "price_summary":
        decision_rows = [
            ("投标总报价（小写）", placeholder("投标总报价（小写）")),
            ("投标总报价（大写）", placeholder("投标总报价（大写）")),
            ("暂列金额", placeholder("暂列金额")),
            ("安全文明施工费", placeholder("安全文明施工费")),
            ("税率", placeholder("税率")),
            ("工期", placeholder("投标工期")),
            ("质量目标", placeholder("质量目标")),
            ("投标有效期", placeholder("投标有效期")),
            ("主要材料价格风险范围", placeholder("主要材料价格风险范围")),
        ]
        rows = [[index, label, value] for index, (label, value) in enumerate(decision_rows, start=1)]
        return "**附表：报价与工期决策字段**\n\n" + _md_table(
            ("序号", "填写项", "内容"), rows
        )
    return ""


def render_profile_tables(profile: BidderProfile, table_ids: list[str]) -> str:
    blocks = []
    for table_id in table_ids:
        table = render_profile_table(profile, table_id)
        if table:
            blocks.append(table)
    return "\n\n".join(blocks)


def _strip_leading_duplicate_heading(lines: list[str], title: str) -> list[str]:
    title_key = normalize_ws(title)
    for index, line in enumerate(lines):
        if not line.strip():
            continue
        match = _HEADING_RE.match(line.strip())
        if match:
            heading = normalize_ws(match.group(2))
            if title_key in heading or heading in title_key:
                return [*lines[:index], *lines[index + 1 :]]
        return lines
    return lines


def _demote_headings(lines: list[str]) -> list[str]:
    result: list[str] = []
    for line in lines:
        match = _HEADING_RE.match(line.strip())
        if not match:
            result.append(line)
            continue
        level = min(len(match.group(1)) + 1, 4)
        result.append(f"{'#' * level} {match.group(2)}")
    return result


def normalize_section_markdown(section: BidSection, draft: str, number: int) -> str:
    text = (draft or "").strip()
    if text.startswith("```"):
        lines = text.splitlines()
        if lines and lines[0].startswith("```"):
            lines = lines[1:]
        if lines and lines[-1].strip().startswith("```"):
            lines = lines[:-1]
        text = "\n".join(lines).strip()
    if not text:
        text = f"【待核对：章节“{section.title}”未生成，请重试或人工补写】"
    lines = _strip_leading_duplicate_heading(text.splitlines(), section.title)
    lines = _demote_headings(lines)
    body = "\n".join(lines).strip()
    heading = f"## {cn_index(number)}、{section.title}"
    return f"{heading}\n\n{body}".strip()


def build_bid_markdown(
    *,
    tender: TenderContext,
    profile: BidderProfile,
    outline: BidOutline,
    drafts: dict[str, str],
    matrix: list[dict[str, Any]],
) -> str:
    project = tender.project
    cover = "\n".join(
        [
            "<!-- cover -->",
            f"# {project.get('project_name', tender.project_name)}",
            "## 投标文件",
            "",
            f"**标段名称：** {project.get('bid_section', '')}",
            f"**招标编号：** {project.get('tender_number', '')}",
            f"**招标人：** {_tender_owner(tender)}",
            f"**投标人：** {profile.company.get('name', '')}",
            f"**投标日期：** {placeholder('投标日期')}",
            "<!-- pagebreak -->",
        ]
    )
    toc_lines = ["## 目录", "", "- 评标要素响应索引表"]
    for index, section in enumerate(outline.sections, start=1):
        toc_lines.append(f"- {cn_index(index)}、{section.title}")
    toc_lines.append("<!-- pagebreak -->")
    toc = "\n".join(toc_lines)

    index_lines = ["## 评标要素响应索引表", ""]
    if matrix:
        rows = []
        for index, row in enumerate(matrix, start=1):
            score = row.get("max_score")
            score_text = f"{score:g}" if isinstance(score, (int, float)) else "-"
            rows.append(
                [
                    index,
                    row.get("factor", ""),
                    row.get("category", ""),
                    score_text,
                    row.get("display", ""),
                ]
            )
        index_lines.append(
            _md_table(("序号", "评标要素", "类别", "分值", "响应章节"), rows)
        )
        index_lines.append("")
        index_lines.append(
            "说明：本表仅用于索引响应章节，分值和评审标准以招标文件为准；"
            "投标人应按评标办法逐项核对，不得将本表作为评标依据。"
        )
    else:
        index_lines.append("【待核对：未从招标文件抽取到评分因素，请人工编制评标要素索引表】")
    index = "\n".join(index_lines)

    parts = [cover, toc, index]
    for number, section in enumerate(outline.sections, start=1):
        body = normalize_section_markdown(
            section, drafts.get(section.id, ""), number
        )
        tables = render_profile_tables(profile, section.tables)
        if tables:
            body = f"{body}\n\n{tables}"
        parts.append("<!-- pagebreak -->")
        parts.append(body)
    return "\n\n".join(parts).strip() + "\n"


def _tender_owner(tender: TenderContext) -> str:
    """The owner is not always extracted as a scalar; use a placeholder."""
    for item in tender.other_requirements:
        match = re.search(r"招标人(?:名称)?[为：:]\s*([^\s；;，,。]{2,60})", item)
        if match:
            return match.group(1)
    section = tender.find_section("联系方式")
    if section:
        match = re.search(r"招标人[：:]\s*([^\n]{2,60})", section.text)
        if match:
            return match.group(1).strip()
    return "【待核对：招标人名称】"


def collect_pending_markers(markdown: str) -> list[PendingMarker]:
    found: dict[str, list[str]] = {}
    order: list[str] = []
    section = "封面或索引"
    for line in markdown.splitlines():
        stripped = line.strip()
        if stripped.startswith("## "):
            section = stripped[3:].strip()
        for match in PENDING_RE.finditer(line):
            marker = match.group(0)
            if marker not in found:
                found[marker] = []
                order.append(marker)
            if section not in found[marker]:
                found[marker].append(section)
    return [PendingMarker(marker=marker, sections=found[marker]) for marker in order]


def _severity_label(severity: str) -> str:
    return {"high": "高", "medium": "中", "low": "低"}.get(severity, severity)


def build_review_markdown(
    *,
    tender: TenderContext,
    profile: BidderProfile,
    outline: BidOutline,
    matrix: list[dict[str, Any]],
    review: dict[str, Any],
    revisions: list[dict[str, Any]],
    warnings: list[str],
    llm_info: dict[str, Any],
    markdown: str,
    generated_at: str | None = None,
) -> str:
    generated_at = generated_at or now_iso()
    pending = collect_pending_markers(markdown)
    issues = review.get("issues") or []
    claims = review.get("unverified_claims") or []
    missing = review.get("missing_requirements") or []
    lines: list[str] = [
        "# 投标文件编制自检报告（内部）",
        "",
        "> 本报告仅供内部编制与复核使用，不随投标文件提交。生成稿中的"
        "【待填写：...】和【待核对：...】必须由人工补齐或核对后才能提交。",
        "",
        "## 1. 生成信息",
        "",
        _md_table(
            ("项目", "内容"),
            [
                ("项目名称", tender.project.get("project_name", "")),
                ("招标编号", tender.project.get("tender_number", "")),
                ("标段名称", tender.project.get("bid_section", "")),
                ("投标人", profile.company.get("name", "")),
                ("源文件", tender.source_name),
                ("解析时间", tender.parsed_at),
                ("生成时间", generated_at),
                ("LLM", f"{llm_info.get('provider', '')}/{llm_info.get('model', '')}"),
                ("LLM 调用次数", str(llm_info.get("calls", 0))),
                ("章节数", str(len(outline.sections))),
                ("审查问题", f"高 {sum(1 for item in issues if item.get('severity') == 'high')} / "
                 f"中 {sum(1 for item in issues if item.get('severity') == 'medium')} / "
                 f"低 {sum(1 for item in issues if item.get('severity') == 'low')}"),
                ("待填写/待核对标记", str(len(pending))),
            ],
        ),
        "",
        "## 2. 审查结论",
        "",
        review.get("summary") or "（审查 Agent 未返回结论，请以问题清单为准）",
        "",
        "## 3. 待填写与待核对清单",
        "",
    ]
    if pending:
        lines.append(
            _md_table(
                ("序号", "标记", "出现章节"),
                [
                    (index, item.marker, "；".join(item.sections))
                    for index, item in enumerate(pending, start=1)
                ],
            )
        )
    else:
        lines.append("未发现占位标记。仍需人工核对所有事实性内容。")

    lines.extend(["", "## 4. 投标人资料缺口", ""])
    if profile.gaps:
        lines.append(
            _md_table(
                ("序号", "资料字段", "占位符", "建议章节", "JSON 路径"),
                [
                    (index, gap.label, gap.placeholder, gap.section, gap.path)
                    for index, gap in enumerate(profile.gaps, start=1)
                ],
            )
        )
    else:
        lines.append("投标人资料未发现缺失字段。")

    lines.extend(["", "## 5. 审查问题", ""])
    if issues:
        for severity, label in (("high", "高"), ("medium", "中"), ("low", "低")):
            grouped = [item for item in issues if item.get("severity") == severity]
            if not grouped:
                continue
            lines.append(f"### {label}优先级问题")
            lines.append("")
            lines.append(
                _md_table(
                    ("序号", "章节", "问题", "修改建议", "招标文件依据"),
                    [
                        (
                            index,
                            item.get("section_id") or "全局",
                            item.get("problem", ""),
                            item.get("fix", ""),
                            item.get("tender_basis", ""),
                        )
                        for index, item in enumerate(grouped, start=1)
                    ],
                )
            )
            lines.append("")
    else:
        lines.append("审查未发现需要记录的问题。")

    lines.extend(["", "## 6. 未经投标人资料支持的事实声明", ""])
    if claims:
        lines.append(
            _md_table(
                ("序号", "章节", "声明内容", "原因"),
                [
                    (
                        index,
                        item.get("section_id") or "未知",
                        item.get("claim", ""),
                        item.get("reason", ""),
                    )
                    for index, item in enumerate(claims, start=1)
                ],
            )
        )
    else:
        lines.append("审查 Agent 未标记额外的无依据事实；仍须人工逐项核对。")

    lines.extend(["", "## 7. 缺失或覆盖不足的招标要求", ""])
    if missing:
        lines.append(
            _md_table(
                ("序号", "要求", "建议章节", "补充建议"),
                [
                    (
                        index,
                        item.get("requirement", ""),
                        item.get("section_id") or "未指定",
                        item.get("suggestion", ""),
                    )
                    for index, item in enumerate(missing, start=1)
                ],
            )
        )
    else:
        lines.append("审查 Agent 未返回缺失要求清单；请结合招标文件逐项复核。")

    lines.extend(["", "## 8. 评分响应矩阵", ""])
    if matrix:
        lines.append(
            _md_table(
                ("序号", "评分因素", "类别", "分值", "响应章节", "评分标准摘要"),
                [
                    (
                        index,
                        item.get("factor", ""),
                        item.get("category", ""),
                        item.get("max_score", ""),
                        item.get("display", ""),
                        str(item.get("criteria", ""))[:240],
                    )
                    for index, item in enumerate(matrix, start=1)
                ],
            )
        )
    else:
        lines.append("未抽取到评分因素。")

    lines.extend(["", "## 9. 修订记录", ""])
    if revisions:
        lines.append(
            _md_table(
                ("章节", "修订前字数", "修订后字数", "针对问题数"),
                [
                    (
                        item.get("section_title", item.get("section_id", "")),
                        item.get("before_chars", 0),
                        item.get("after_chars", 0),
                        len(item.get("issues") or []),
                    )
                    for item in revisions
                ],
            )
        )
    else:
        lines.append("未执行自动修订。")

    lines.extend(["", "## 10. 生成警告与数据来源", ""])
    if warnings:
        lines.extend(f"- {warning}" for warning in warnings)
    else:
        lines.append("- 无生成警告。")
    lines.extend(
        [
            "",
            f"- 解析清单：`{tender.manifest.get('artifacts_dir', '')}/manifest.json`",
            f"- 结构化抽取：`{tender.manifest.get('artifacts_dir', '')}/extraction.json`",
            f"- 投标人资料：`{profile.source_path or '未提供'}`",
            f"- 摘要哈希：`{profile.sha256 or '未提供'}`",
            "",
            "## 11. 提交前人工检查",
            "",
            "- [ ] 补齐全部【待填写：...】并核对全部【待核对：...】",
            "- [ ] 核对投标总报价、工期、质量目标、投标有效期和保证金，确保不是招标估算价",
            "- [ ] 核对资质、证书编号、有效期、人员社保及证书扫描件",
            "- [ ] 核对业绩的项目名称、金额、日期、角色和证明材料",
            "- [ ] 按招标文件第八章复核格式、签字、盖章和电子印章要求",
            "- [ ] 核对投标文件格式、上传方式、投标截止时间和开标安排",
            "- [ ] 删除本内部报告，不将其编入投标文件",
        ]
    )
    return "\n".join(lines).strip() + "\n"


# ---------------------------------------------------------------------------
# Word export
# ---------------------------------------------------------------------------


def _set_run_font(run: Any, *, east_asia: str, size: float, bold: bool = False, color: str | None = None) -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), east_asia)


def _set_style_font(style: Any, *, east_asia: str, size: float, bold: bool = False, color: str | None = None) -> None:
    style.font.name = "Times New Roman"
    style.font.size = Pt(size)
    style.font.bold = bold
    if color:
        style.font.color.rgb = RGBColor.from_string(color)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), east_asia)


def _add_page_number_footer(section: Any) -> None:
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instruction)
    run._r.append(separate)
    run._r.append(text)
    run._r.append(end)
    _set_run_font(run, east_asia="宋体", size=9)


def _add_header(section: Any, text: str) -> None:
    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run(text)
    _set_run_font(run, east_asia="宋体", size=9, color="666666")


def _style_document(document: Document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    _set_style_font(normal, east_asia="宋体", size=11)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.first_line_indent = Pt(22)

    for name, size, color in (
        ("Heading 1", 18, "1F3864"),
        ("Heading 2", 16, "1F3864"),
        ("Heading 3", 14, "1F3864"),
        ("Heading 4", 12, "1F3864"),
    ):
        style = styles[name]
        _set_style_font(style, east_asia="黑体", size=size, bold=True, color=color)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        _set_style_font(style, east_asia="宋体", size=11)
        style.paragraph_format.first_line_indent = Pt(0)


def _split_table_row(line: str) -> list[str]:
    text = line.strip()
    if text.startswith("|"):
        text = text[1:]
    if text.endswith("|"):
        text = text[:-1]
    sentinel = "\x00PIPE\x00"
    text = text.replace("\\|", sentinel)
    return [cell.strip().replace(sentinel, "|") for cell in text.split("|")]


def _is_table_separator(line: str) -> bool:
    if "|" not in line:
        return False
    cells = [cell.strip() for cell in _split_table_row(line)]
    return bool(cells) and all(_TABLE_SEPARATOR_RE.match(cell) for cell in cells if cell)


def _add_inline(paragraph: Any, text: str, *, size: float = 11, east_asia: str = "宋体") -> None:
    position = 0
    for match in _INLINE_RE.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position : match.start()])
            _set_run_font(run, east_asia=east_asia, size=size)
        token = match.group(0)
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            _set_run_font(run, east_asia=east_asia, size=size, bold=True)
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(size - 0.5)
        elif token.startswith("["):
            label = token[1 : token.index("]")]
            run = paragraph.add_run(label)
            _set_run_font(run, east_asia=east_asia, size=size)
            run.font.underline = True
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        _set_run_font(run, east_asia=east_asia, size=size)


def _add_body_paragraph(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="Normal")
    _add_inline(paragraph, text)


def _add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    columns = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=columns)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for row_index, row in enumerate(rows):
        for column_index in range(columns):
            cell = table.cell(row_index, column_index)
            cell.text = ""
            value = row[column_index] if column_index < len(row) else ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.first_line_indent = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.15
            _add_inline(paragraph, value, size=9)
            if row_index == 0:
                for run in paragraph.runs:
                    run.font.bold = True
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "D9E2F3")
                cell._tc.get_or_add_tcPr().append(shading)
    document.add_paragraph()


def _add_cover(document: Document, lines: list[str], project_text: str) -> None:
    for _ in range(4):
        document.add_paragraph()
    for line in lines:
        stripped = line.strip()
        if not stripped:
            document.add_paragraph()
            continue
        if stripped.startswith("# "):
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(stripped[2:])
            _set_run_font(run, east_asia="黑体", size=24, bold=True)
        elif stripped.startswith("## "):
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(stripped[3:])
            _set_run_font(run, east_asia="黑体", size=20, bold=True, color="1F3864")
        else:
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(10)
            _add_inline(paragraph, stripped, size=13)
    document.add_paragraph()
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    count = len(collect_pending_markers("\n".join(lines)))
    if count:
        run = paragraph.add_run(f"（草稿：含 {count} 类待填写事项，提交前请逐项核实）")
        _set_run_font(run, east_asia="楷体", size=10, color="C00000")


def export_docx(
    markdown: str,
    path: str | Path,
    *,
    header_text: str = "",
) -> Path:
    """Convert the generated Markdown to a styled Word document."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.4)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.6)
    section.right_margin = Cm(2.2)
    section.different_first_page_header_footer = True
    _style_document(document)
    if header_text:
        _add_header(section, header_text)
    _add_page_number_footer(section)

    lines = markdown.replace("\r\n", "\n").splitlines()
    index = 0
    in_cover = False
    cover_lines: list[str] = []
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if stripped == "<!-- cover -->":
            in_cover = True
            cover_lines = []
            index += 1
            continue
        if in_cover:
            if stripped == "<!-- pagebreak -->":
                _add_cover(document, cover_lines, header_text)
                document.add_page_break()
                in_cover = False
            else:
                cover_lines.append(line)
            index += 1
            continue
        if stripped == "<!-- pagebreak -->":
            document.add_page_break()
            index += 1
            continue
        if stripped.startswith("<!--") and stripped.endswith("-->"):
            index += 1
            continue
        if not stripped:
            index += 1
            continue
        heading = _HEADING_RE.match(stripped)
        if heading:
            level = len(heading.group(1))
            paragraph = document.add_heading(level=min(level, 4))
            _add_inline(
                paragraph,
                heading.group(2),
                size={1: 18, 2: 16, 3: 14, 4: 12}.get(level, 11),
                east_asia="黑体",
            )
            index += 1
            continue
        if stripped.startswith("|") and index + 1 < len(lines) and _is_table_separator(lines[index + 1]):
            table_rows: list[list[str]] = [_split_table_row(stripped)]
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_rows.append(_split_table_row(lines[index]))
                index += 1
            _add_table(document, table_rows)
            continue
        bullet = _BULLET_RE.match(stripped)
        if bullet:
            content = bullet.group(1).strip()
            if content.startswith("[ ]"):
                content = "☐ " + content[3:].strip()
            paragraph = document.add_paragraph(style="List Bullet")
            _add_inline(paragraph, content)
            index += 1
            continue
        numbered = _NUMBERED_RE.match(stripped)
        if numbered:
            paragraph = document.add_paragraph(style="List Number")
            _add_inline(paragraph, numbered.group(2))
            index += 1
            continue
        if stripped.startswith(">"):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.6)
            paragraph.paragraph_format.first_line_indent = Pt(0)
            run = paragraph.add_run(stripped.lstrip("> ").strip())
            _set_run_font(run, east_asia="楷体", size=10, color="555555")
            run.font.italic = True
            index += 1
            continue
        if stripped in {"---", "***", "___"}:
            index += 1
            continue
        # Plain paragraph: merge following non-block lines.
        paragraph_lines = [stripped]
        index += 1
        while index < len(lines):
            candidate = lines[index].strip()
            if (
                not candidate
                or candidate.startswith("#")
                or candidate.startswith("|")
                or candidate.startswith(">")
                or candidate.startswith("<!--")
                or _BULLET_RE.match(candidate)
                or _NUMBERED_RE.match(candidate)
                or candidate in {"---", "***", "___"}
            ):
                break
            paragraph_lines.append(candidate)
            index += 1
        _add_body_paragraph(document, " ".join(paragraph_lines))

    if in_cover and cover_lines:
        _add_cover(document, cover_lines, header_text)

    document.core_properties.title = header_text or path.stem
    document.core_properties.comments = "由 bidding 多 Agent 流程生成的草稿，含待填写占位符。"
    temporary = path.with_suffix(path.suffix + ".tmp")
    document.save(str(temporary))
    temporary.replace(path)
    return path
