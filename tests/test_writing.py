"""Offline tests for the multi-agent bid-writing pipeline."""

from __future__ import annotations

import json
from pathlib import Path
from types import SimpleNamespace
from typing import Any

import pytest
from docx import Document

from bidding.cli import main
from bidding.writing.agents import deterministic_review
from bidding.writing.errors import BidWriteError
from bidding.writing.graph import run_writing_graph
from bidding.writing.llm import OpenAIBidLLM
from bidding.writing.outline import (
    build_outline,
    build_response_matrix,
    canonical_sections,
)
from bidding.writing.pipeline import write_bid
from bidding.writing.profile import (
    has_placeholder,
    load_profile,
    normalize_profile,
)
from bidding.writing.render import (
    build_review_markdown,
    cn_index,
    collect_pending_markers,
    normalize_section_markdown,
)
from bidding.writing.retrieval import TenderRetriever, format_evidence
from bidding.writing.tender import load_tender_context
from support import make_settings

DOC_ID = "doc-write"


def _write_json(path: Path, data: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def _fixture(tmp_path: Path):
    settings = make_settings(tmp_path)
    settings.ensure_directories()
    parsed = settings.parsed_dir / DOC_ID
    sections = [
        {
            "index": 1,
            "level": 1,
            "number": "第二章",
            "title": "第二章 投标人须知",
            "path": ["第二章 投标人须知"],
            "page_start": 2,
            "page_end": 2,
            "text": (
                "<!-- page: 2 -->\n第二章 投标人须知\n"
                "投标人须具有水利水电工程施工总承包二级及以上资质。"
                "投标文件应按照第八章格式编制，并加盖电子印章。"
            ),
        },
        {
            "index": 2,
            "level": 1,
            "number": "第三章",
            "title": "第三章 评标办法（综合评估法）",
            "path": ["第三章 评标办法（综合评估法）"],
            "page_start": 3,
            "page_end": 3,
            "text": (
                "<!-- page: 3 -->\n第三章 评标办法\n"
                "本项目采用综合评估法。施工方案与技术措施，满分 40 分。"
                "质量管理体系与措施，满分 10 分。投标报价，满分 30 分。"
                "安全管理体系与措施，满分 5 分。环境保护管理体系与措施，满分 5 分。"
                "施工方案应包含施工总部署、关键工序和质量安全措施。"
            ),
        },
        {
            "index": 3,
            "level": 1,
            "number": "第八章",
            "title": "第八章 投标文件格式",
            "path": ["第八章 投标文件格式"],
            "page_start": 5,
            "page_end": 5,
            "text": (
                "<!-- page: 5 -->\n第八章 投标文件格式\n"
                "投标函、法定代表人身份证明、授权委托书、投标保证金等格式详见本章。"
            ),
        },
    ]
    _write_json(
        parsed / "manifest.json",
        {
            "doc_id": DOC_ID,
            "source_name": "tender.pdf",
            "parsed_at": "2026-10-19T10:00:00+08:00",
            "artifacts_dir": str(parsed),
            "warnings": [],
        },
    )
    _write_json(
        parsed / "raw" / "sections.json",
        {"sections": sections, "toc": []},
    )
    extraction = {
        "project_name": "测试供水工程",
        "tender_number": "T-2026-001",
        "bid_section": "测试供水工程（施工）",
        "project_scale": "新建供水管网及附属构筑物",
        "budget_amount_text": "333.37万元",
        "bid_deadline": "2026-11-14 13:30",
        "bid_opening_time": "2026-11-14 13:30",
        "bid_opening_place": "测试开标室",
        "evaluation_method": "综合评估法",
        "evaluation_factors": [
            {
                "factor": "施工方案与技术措施",
                "category": "技术",
                "max_score": 40,
                "criteria": "施工方案完整、针对性强，重点难点措施可行。",
                "source_quote": "施工方案与技术措施，满分 40 分",
            },
            {
                "factor": "质量管理体系与措施",
                "category": "技术",
                "max_score": 10,
                "criteria": "质量管理体系健全，措施具体。",
                "source_quote": "质量管理体系与措施，满分 10 分",
            },
            {
                "factor": "安全管理体系与措施",
                "category": "技术",
                "max_score": 5,
                "criteria": "安全管理体系完整，措施得力。",
                "source_quote": "安全管理体系与措施，满分 5 分",
            },
            {
                "factor": "环境保护管理体系与措施",
                "category": "技术",
                "max_score": 5,
                "criteria": "环境保护管理体系完整，措施得力。",
                "source_quote": "环境保护管理体系与措施，满分 5 分",
            },
            {
                "factor": "投标报价",
                "category": "报价",
                "max_score": 30,
                "criteria": "报价合理。",
                "source_quote": "投标报价，满分 30 分",
            },
        ],
        "preliminary_review": [
            {
                "stage": "资格性审查",
                "item": "营业执照",
                "criteria": "营业执照有效",
                "source_quote": "投标人须具有水利水电工程施工总承包二级及以上资质",
            }
        ],
        "qualifications": [
            {
                "category": "资质",
                "requirement": "具有水利水电工程施工总承包二级及以上资质",
                "thresholds": ["二级"],
                "source_quote": "投标人须具有水利水电工程施工总承包二级及以上资质",
            }
        ],
        "bid_format": {
            "items": [
                {
                    "title": "投标函",
                    "requirements": "按第八章格式填写并加盖电子印章",
                }
            ]
        },
        "technical_requirements": [
            {
                "topic": "施工组织设计",
                "requirement": "施工方案应包含施工总部署、关键工序和质量安全措施。",
                "source_quote": "施工方案应包含施工总部署、关键工序和质量安全措施",
            }
        ],
        "contract_points": [
            {
                "topic": "工期",
                "detail": "工期以合同约定为准。",
                "source_quote": "",
            }
        ],
        "other_requirements": ["投标文件应加盖电子印章。"],
        "key_dates": [
            {"event": "投标截止时间", "time": "2026-11-14 13:30"}
        ],
        "ungrounded_fields": [],
        "extraction_meta": {
            "partial": False,
            "failed_units": [],
            "processed_units": 1,
            "unit_count": 1,
        },
    }
    _write_json(parsed / "extraction.json", extraction)
    profile_path = tmp_path / "bidder_profile.json"
    _write_json(
        profile_path,
        {
            "company": {
                "name": "示例建设有限公司",
                "unified_social_credit_code": "",
                "legal_representative": "示例法定代表人",
                "registered_capital": "",
                "established_date": "",
                "address": "",
                "contact_person": "",
                "phone": "",
                "email": "",
            },
            "qualifications": [
                {
                    "name": "水利水电工程施工总承包",
                    "level": "二级",
                    "certificate_no": "TEST-QUAL-001",
                    "issuer": "",
                    "valid_until": "2028-12-31",
                    "note": "",
                }
            ],
            "similar_projects": [
                {
                    "project_name": "示例供水工程",
                    "owner": "示例招标人",
                    "contract_amount": "500万元",
                    "start_date": "2022-01-01",
                    "end_date": "2022-12-31",
                    "role": "施工总承包",
                    "description": "",
                    "proof": "",
                }
            ],
            "personnel": [
                {
                    "name": "示例项目经理",
                    "role": "项目经理",
                    "title": "工程师",
                    "certificate": "二级建造师",
                    "certificate_no": "TEST-PM-001",
                    "experience_years": "10",
                    "projects": "",
                    "note": "",
                }
            ],
            "financials": [
                {
                    "year": "2022",
                    "revenue": "1000万元",
                    "net_profit": "100万元",
                    "total_assets": "2000万元",
                    "total_liabilities": "800万元",
                    "audit_status": "已审计",
                }
            ],
            "credit": {},
            "management_systems": [],
            "safety": {},
            "commitments": ["按招标文件要求完成工程"],
            "capabilities": ["具备类似工程施工经验"],
            "notes": "",
        },
    )
    return settings, profile_path


class FakeBidLLM:
    provider = "fake"
    model = "fake-writer"
    records: list[Any] = []

    def __init__(self) -> None:
        self.calls = 0
        self.purposes: list[str] = []
        self.user_prompts: list[str] = []

    @staticmethod
    def _title(user: str) -> str:
        marker = "- 章节："
        for line in user.splitlines():
            if line.strip().startswith(marker):
                return line.split(marker, 1)[1].strip()
        return "测试章节"

    def complete_json(self, system: str, user: str, *, purpose: str = "") -> dict[str, Any]:
        self.calls += 1
        self.purposes.append(purpose)
        if purpose == "analyst":
            return {
                "project_summary": "本项目为供水工程施工，重点响应施工方案与质量要求。",
                "evaluation_strategy": [
                    {
                        "factor": "施工方案与技术措施",
                        "section_id": "technical_methods",
                        "response_approach": "按工序逐项响应。",
                    }
                ],
                "section_briefs": {
                    "technical_methods": "围绕供水管网和构筑物展开，覆盖关键工序。"
                },
                "extra_sections": [],
                "risks": ["投标报价需投标人自主决策，不得沿用估算价。"],
            }
        if purpose == "reviewer":
            return {
                "summary": "草稿结构完整，但存在预算金额误用风险，需修订。",
                "issues": [],
                "unverified_claims": [],
                "missing_requirements": [],
            }
        raise AssertionError(f"unexpected JSON purpose: {purpose}")

    def complete_text(self, system: str, user: str, *, purpose: str = "") -> str:
        self.calls += 1
        self.purposes.append(purpose)
        self.user_prompts.append(user)
        title = self._title(user)
        if purpose.startswith("reviser:"):
            return (
                f"## {title}\n\n已按审查意见修订。投标报价为【待填写：投标总报价】，"
                "不以招标估算价代替。"
            )
        if purpose.startswith("writer:"):
            section_id = purpose.split(":", 1)[1]
            if section_id == "price_commitments":
                return (
                    f"## {title}\n\n本项目合同估算价为 333.37 万元，"
                    "投标报价为【待填写：投标总报价】。"
                )
            return (
                f"## {title}\n\n本章围绕测试供水工程编制，"
                "对招标文件要求逐项响应，缺失事实使用【待填写：...】。"
            )
        raise AssertionError(f"unexpected text purpose: {purpose}")


def test_profile_missing_fields_become_placeholders(tmp_path: Path) -> None:
    profile = normalize_profile(
        {
            "company": {"name": "示例公司"},
            "qualifications": [{"name": "资质 A"}],
            "personnel": [],
        }
    )
    assert profile.company["name"] == "示例公司"
    assert has_placeholder(profile.company["legal_representative"])
    assert has_placeholder(profile.company["address"])
    assert len(profile.qualifications) == 1
    assert has_placeholder(profile.qualifications[0]["certificate_no"])
    assert len(profile.personnel) == 1
    assert all(has_placeholder(value) for value in profile.personnel[0].values())
    labels = {gap.label for gap in profile.gaps}
    assert {"统一社会信用代码", "证书编号", "人员姓名", "年度"} <= labels
    assert profile.provided_field_count == 2
    assert "【待填写：统一社会信用代码】" in profile.prompt_text()

    profile_file = tmp_path / "profile.json"
    _write_json(profile_file, {"company": {"name": "示例公司"}})
    loaded = load_profile(profile_file)
    assert loaded.company_name == "示例公司"
    assert loaded.sha256
    with pytest.raises(BidWriteError, match="找不到投标人资料"):
        load_profile(tmp_path / "missing.json")


def test_prompt_trimming_does_not_mutate_profile_facts() -> None:
    long_name = "示例公司" * 200
    profile = normalize_profile({"company": {"name": long_name}})
    text = profile.prompt_text(max_chars=5000)
    assert len(text) <= 5000
    assert profile.company["name"] == long_name
    assert profile.company_name == long_name


def test_bid_llm_fails_fast_on_auth_error() -> None:
    class StatusError(Exception):
        status_code = 401

    llm = object.__new__(OpenAIBidLLM)
    llm.settings = SimpleNamespace(llm_max_retries=3, mineru_retry_backoff=1.0)
    llm.provider = "zhipu"
    llm.records = []
    llm._create = lambda *args, **kwargs: (_ for _ in ()).throw(
        StatusError("invalid api key")
    )
    with pytest.raises(BidWriteError) as excinfo:
        llm.complete_text("系统", "用户", purpose="writer:test")
    assert excinfo.value.fatal
    assert llm.calls == 0


def test_tender_context_and_retrieval_keep_page_citations(tmp_path: Path) -> None:
    settings, _profile_path = _fixture(tmp_path)
    tender = load_tender_context(DOC_ID, settings)
    assert tender.project_name == "测试供水工程"
    assert tender.project["budget_amount_text"] == "333.37万元"
    assert tender.locate_quote("施工方案与技术措施，满分 40 分") == (3, 3)

    outline = build_outline({}, tender)
    assert len(outline.sections) == 11
    retriever = TenderRetriever(tender)
    evidence, warnings = retriever.retrieve(outline.by_id()["technical_methods"])
    assert not warnings
    assert any(item.kind == "评分因素" for item in evidence)
    original = [item for item in evidence if item.kind == "招标文件原文"]
    assert original
    assert original[0].page_start == 3
    prompt = format_evidence(evidence)
    assert "施工方案与技术措施" in prompt
    assert "第3页" in prompt


def test_outline_is_deterministic_and_accepts_bounded_analyst_input(tmp_path: Path) -> None:
    settings, _profile_path = _fixture(tmp_path)
    tender = load_tender_context(DOC_ID, settings)
    analysis = {
        "project_summary": "项目理解",
        "evaluation_strategy": [
            {
                "factor": "施工方案与技术措施",
                "section_id": "technical_methods",
                "response_approach": "逐项响应",
            }
        ],
        "section_briefs": {"technical_methods": "补充重点难点"},
        "extra_sections": [
            {
                "title": "施工组织设计补充说明",
                "role": "technical",
                "purpose": "补充说明",
                "instructions": ["说明接口"],
                "keywords": ["接口"],
                "word_target": 800,
            },
            {
                "title": "施工方案与技术措施",
                "role": "technical",
            },
        ],
        "risks": ["风险一"],
    }
    outline = build_outline(analysis, tender)
    ids = [section.id for section in outline.sections]
    assert len(outline.sections) == 12
    assert "x1" in ids
    assert canonical_sections()[5].id == "technical_overview"
    assert any("分析师提示" in item for item in outline.by_id()["technical_methods"].instructions)
    assert outline.warnings
    assert outline.by_id()["x1"].origin == "analyst"

    matrix = build_response_matrix(outline, tender)
    by_factor = {row["factor"]: row for row in matrix}
    assert by_factor["施工方案与技术措施"]["section_ids"] == ["technical_methods"]
    assert by_factor["质量管理体系与措施"]["section_ids"] == ["quality_management"]
    assert by_factor["安全管理体系与措施"]["section_ids"] == ["safety_environment"]
    assert by_factor["环境保护管理体系与措施"]["section_ids"] == ["safety_environment"]
    assert by_factor["投标报价"]["section_ids"] == ["price_commitments"]


def test_graph_runs_all_agents_and_revises_budget_misuse(tmp_path: Path) -> None:
    settings, profile_path = _fixture(tmp_path)
    profile = load_profile(profile_path)
    tender = load_tender_context(DOC_ID, settings)
    retriever = TenderRetriever(tender)
    llm = FakeBidLLM()
    from bidding.writing.agents import BidAgents

    agents = BidAgents(llm, profile, tender, retriever)
    progress: list[str] = []
    state = run_writing_graph(
        agents=agents,
        retriever=retriever,
        tender=tender,
        profile=profile,
        max_revisions=1,
        progress=progress.append,
    )
    outline = state["outline"]
    assert len(outline.sections) == 11
    assert len(state["drafts"]) == 11
    assert any(
        "投标报价" in issue["problem"] for issue in state["review"]["issues"]
    )
    assert len(state["revisions"]) == 1
    assert state["revisions"][0]["section_id"] == "price_commitments"
    assert "333.37 万元" not in state["drafts"]["price_commitments"]
    assert "【待填写：投标总报价】" in state["drafts"]["price_commitments"]
    assert state["markdown"].startswith("<!-- cover -->")
    assert "## 一、投标函及投标函附录" in state["markdown"]
    assert "**附表：投标人基本情况**" in state["markdown"]
    assert "示例建设有限公司" in state["markdown"]
    assert "333.37 万元" not in state["markdown"]
    node_names = [entry["node"] for entry in state["trace"]]
    assert node_names == [
        "analyst",
        "retriever",
        "admin_writer",
        "commercial_writer",
        "technical_writer",
        "reviewer",
        "reviser",
        "assembler",
    ]
    assert any("正在编写" in message for message in progress)
    technical_prompt = next(
        item for item in llm.user_prompts if "- 章节：施工方案与技术措施" in item
    )
    assert (
        "【本章节必须明确响应的评分因素】\n- 施工方案与技术措施"
        in technical_prompt
    )
    safety_prompt = next(
        item for item in llm.user_prompts if "- 章节：安全管理、环境保护与文明施工" in item
    )
    assert "安全管理体系与措施" in safety_prompt
    assert "环境保护管理体系与措施" in safety_prompt


def test_graph_skips_reviser_when_disabled(tmp_path: Path) -> None:
    settings, profile_path = _fixture(tmp_path)
    profile = load_profile(profile_path)
    tender = load_tender_context(DOC_ID, settings)
    retriever = TenderRetriever(tender)
    llm = FakeBidLLM()
    from bidding.writing.agents import BidAgents

    state = run_writing_graph(
        agents=BidAgents(llm, profile, tender, retriever),
        retriever=retriever,
        tender=tender,
        profile=profile,
        max_revisions=0,
    )
    assert state["revisions"] == []
    assert "reviser" not in [entry["node"] for entry in state["trace"]]
    assert "333.37 万元" in state["drafts"]["price_commitments"]


def test_deterministic_review_flags_system_fill_and_bad_reference(
    tmp_path: Path,
) -> None:
    settings, profile_path = _fixture(tmp_path)
    profile = load_profile(profile_path)
    tender = load_tender_context(DOC_ID, settings)
    outline = build_outline({}, tender)
    drafts = {section.id: "本章内容。" for section in outline.sections}
    drafts["bid_letter"] += (
        "以上【待填写】内容将由系统根据投标人资料自动填充。"
        "详见投标文件“十、原件的扫描件”章节。"
    )
    issues = deterministic_review(tender, outline, drafts, profile, [])
    assert any("系统自动填充" in item["problem"] for item in issues)
    assert any("不存在的章节" in item["problem"] for item in issues)


def test_graph_ignores_placeholder_only_unverified_claims(tmp_path: Path) -> None:
    class ClaimReviewLLM(FakeBidLLM):
        def complete_json(
            self, system: str, user: str, *, purpose: str = ""
        ) -> dict[str, Any]:
            data = super().complete_json(system, user, purpose=purpose)
            if purpose == "reviewer":
                data["unverified_claims"] = [
                    {
                        "section_id": "company_qualification",
                        "claim": "投标人已获得【待填写：体系名称】认证。",
                        "reason": "投标人资料未提供认证信息",
                    },
                    {
                        "section_id": "company_qualification",
                        "claim": "投标人近三年均盈利。",
                        "reason": "投标人资料未提供财务数据",
                    },
                ]
            return data

    settings, profile_path = _fixture(tmp_path)
    profile = load_profile(profile_path)
    tender = load_tender_context(DOC_ID, settings)
    retriever = TenderRetriever(tender)
    from bidding.writing.agents import BidAgents

    state = run_writing_graph(
        agents=BidAgents(ClaimReviewLLM(), profile, tender, retriever),
        retriever=retriever,
        tender=tender,
        profile=profile,
        max_revisions=0,
    )
    claims = state["review"]["unverified_claims"]
    assert [item["claim"] for item in claims] == ["投标人近三年均盈利。"]
    problems = [item["problem"] for item in state["review"]["issues"]]
    assert any("近三年均盈利" in problem for problem in problems)
    assert not any("体系名称" in problem for problem in problems)


def test_write_bid_writes_markdown_docx_review_and_manifest(tmp_path: Path) -> None:
    settings, profile_path = _fixture(tmp_path)
    llm = FakeBidLLM()
    progress: list[str] = []
    result = write_bid(
        DOC_ID,
        profile_path=profile_path,
        output_dir=tmp_path / "out",
        settings=settings,
        llm=llm,
        max_revisions=1,
        progress=progress.append,
    )
    assert result.markdown_path.is_file()
    assert result.docx_path.is_file()
    assert result.review_path.is_file()
    assert result.manifest_path.is_file()

    markdown = result.markdown_path.read_text(encoding="utf-8")
    assert markdown.startswith("<!-- cover -->")
    assert "## 评标要素响应索引表" in markdown
    assert "## 十一、配合协调、防汛度汛与服务承诺" in markdown
    assert "【待填写：投标日期】" in markdown
    assert "333.37 万元" not in markdown

    review = result.review_path.read_text(encoding="utf-8")
    assert "投标文件编制自检报告（内部）" in review
    assert "待填写与待核对清单" in review
    assert "投标报价" in review
    assert "修订记录" in review

    manifest = json.loads(result.manifest_path.read_text(encoding="utf-8"))
    assert manifest["doc_id"] == DOC_ID
    assert manifest["stats"]["sections"] == 11
    assert manifest["stats"]["revisions"] == 1
    assert manifest["stats"]["pending_markers"] > 0
    assert manifest["bidder"]["company_name"] == "示例建设有限公司"
    assert manifest["llm"]["calls"] == llm.calls
    assert manifest["trace"][-1]["node"] == "assembler"
    assert "API_KEY" not in json.dumps(manifest)

    word = Document(result.docx_path)
    text = "\n".join(paragraph.text for paragraph in word.paragraphs)
    assert "投标文件" in text
    assert "投标函及投标函附录" in text
    assert len(word.tables) >= 8
    assert any("示例建设有限公司" in cell.text for table in word.tables for row in table.rows for cell in row.cells)
    assert any("修订" in message for message in progress)


def test_review_markdown_lists_pending_markers_and_warnings(tmp_path: Path) -> None:
    settings, profile_path = _fixture(tmp_path)
    profile = load_profile(profile_path)
    tender = load_tender_context(DOC_ID, settings)
    outline = build_outline({}, tender)
    markdown = "## 目录\n\n【待填写：投标日期】\n\n## 一、投标函\n\n【待核对：开标地点】\n"
    pending = collect_pending_markers(markdown)
    assert pending[0].marker == "【待填写：投标日期】"
    assert pending[1].sections == ["一、投标函"]
    review = build_review_markdown(
        tender=tender,
        profile=profile,
        outline=outline,
        matrix=[],
        review={
            "summary": "测试",
            "issues": [
                {
                    "severity": "medium",
                    "section_id": "company_qualification",
                    "problem": "测试问题",
                    "fix": "测试修改",
                    "tender_basis": "",
                }
            ],
            "unverified_claims": [],
            "missing_requirements": [],
        },
        revisions=[],
        warnings=["测试警告"],
        llm_info={"provider": "fake", "model": "fake", "calls": 1},
        markdown=markdown,
    )
    assert "【待填写：投标日期】" in review
    assert "测试警告" in review
    assert "未执行自动修订" in review
    assert "投标人基本情况、资质与信用" in review
    assert "company_qualification" not in review


def test_markdown_helpers_normalize_headings_and_numbering() -> None:
    section = canonical_sections()[0]
    text = normalize_section_markdown(
        section,
        f"## {section.title}\n\n### 小节\n正文【待填写：...】【待核对：…】",
        11,
    )
    assert text.startswith("## 十一、投标函及投标函附录")
    assert text.count(section.title) == 1
    assert "#### 小节" in text
    assert "【待填写：...】" not in text
    assert "【待填写】" in text
    assert "【待核对】" in text
    assert cn_index(1) == "一"
    assert cn_index(10) == "十"
    assert cn_index(11) == "十一"
    assert cn_index(20) == "二十"


def test_write_error_when_extraction_missing(tmp_path: Path) -> None:
    settings, profile_path = _fixture(tmp_path)
    (settings.parsed_dir / DOC_ID / "extraction.json").unlink()
    with pytest.raises(BidWriteError, match="尚未完成结构化抽取"):
        write_bid(
            DOC_ID,
            profile_path=profile_path,
            settings=settings,
            llm=FakeBidLLM(),
        )


def test_cli_write_prints_artifact_paths(monkeypatch, tmp_path: Path, capsys) -> None:
    class Result:
        doc_id = DOC_ID
        markdown_path = tmp_path / "bid.md"
        docx_path = tmp_path / "bid.docx"
        review_path = tmp_path / "bid_review.md"
        manifest_path = tmp_path / "bid_manifest.json"
        manifest = {
            "project": {"project_name": "测试供水工程"},
            "bidder": {"company_name": "示例建设有限公司"},
            "stats": {
                "sections": 11,
                "draft_chars": 100,
                "revisions": 1,
                "issues": 2,
                "high_issues": 1,
                "medium_issues": 1,
                "low_issues": 0,
                "pending_markers": 3,
            },
            "llm": {"calls": 14},
            "warnings": ["测试警告"],
        }

    captured: dict[str, Any] = {}

    def fake_write_bid(*args, **kwargs):
        captured.update(kwargs)
        return Result()

    monkeypatch.setattr("bidding.cli.write_bid", fake_write_bid)
    assert main(["write", DOC_ID, "--no-revise", "--max-revisions", "0"]) == 0
    output = capsys.readouterr().out
    assert "测试供水工程" in output
    assert "示例建设有限公司" in output
    assert "sections      : 11" in output
    assert "测试警告" in output
    assert captured["enable_revision"] is False
    assert captured["max_revisions"] == 0

    assert main(["write", DOC_ID, "--max-revisions", "-1"]) == 1
    error = capsys.readouterr().err
    assert "不能为负数" in error
